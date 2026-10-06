#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Ultimate One-Click File Generator for Windows (80+ Formats - Complete Final)
- Fully restores .doc/.xls/.ppt with stable COM DispatchEx
- Generates valid TTF by copying system font
- Generates DXF with actual geometric shapes (line + circle)
- Keeps all newly added text formats
- Includes auto-cleanup prompt and detailed final report
"""

import os
import sys
import csv
import json
import zipfile
import gzip
import bz2
import lzma
import tarfile
import shutil
import subprocess
import io
import wave
import struct
import math
import sqlite3
import glob

# Third-party libraries
try:
    import docx
    from openpyxl import Workbook
    from pptx import Presentation
    from PIL import Image, ImageDraw
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    from odf.opendocument import OpenDocumentText, OpenDocumentSpreadsheet, OpenDocumentPresentation
    from odf.text import P, H
    from odf.table import Table, TableRow, TableCell
    from odf.draw import Page, Frame, TextBox
    from ebooklib import epub
    import fontTools
    from fontTools.fontBuilder import FontBuilder
    import fontTools.pens.ttGlyphPen
    import pycdlib
    import py7zr
    import win32com.client
    import pythoncom
    HAS_ALL_LIBS = True
except ImportError as e:
    HAS_ALL_LIBS = False
    MISSING_LIB = str(e)

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sample_files")

D_DOC = "01_documents"
D_SHEET = "02_sheets"
D_IMG = "03_images"
D_AUD = "04_audio"
D_VID = "05_video"
D_ARC = "06_archives"
D_APP = "07_programs"
D_CODE = "08_code_data"
D_OTH = "09_other"

OK = []
SKIPPED = []
FAILED = []

def ok(name): OK.append(name)
def skip(name, why): SKIPPED.append((name, why))
def fail(name, err): FAILED.append((name, str(err)))

def ensure_dir(path):
    d = os.path.dirname(path)
    if d and not os.path.isdir(d):
        os.makedirs(d)

def find_tool(name, paths):
    p = shutil.which(name)
    if p: return p
    for path in paths:
        if os.path.isfile(path):
            return path
    return None

# --- Find External Tools ---
def find_rar():
    drives = ["C:", "D:", "E:", "F:"]
    paths = [r"Program Files\WinRAR\Rar.exe", r"Program Files (x86)\WinRAR\Rar.exe"]
    p = shutil.which("rar")
    if p: return p
    for d in drives:
        for sub in paths:
            full_path = os.path.join(d + "\\", sub)
            if os.path.isfile(full_path): return full_path
    try:
        import winreg
        key = winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\WinRAR", 0, winreg.KEY_READ)
        val, _ = winreg.QueryValueEx(key, "exe64")
        if val and os.path.isfile(val): return val
    except: pass
    return None

FFMPEG = find_tool("ffmpeg", [r"C:\ffmpeg\bin\ffmpeg.exe", r"D:\ffmpeg\bin\ffmpeg.exe"])
RAR_EXE = find_rar()
CALIBRE_EXE = find_tool("ebook-convert", [r"C:\Program Files\Calibre2\ebook-convert.exe", r"D:\Program Files\Calibre2\ebook-convert.exe"])
CSC_EXE = None
windir = os.environ.get("WINDIR", r"C:\Windows")
for c in (os.path.join(windir, "Microsoft.NET", "Framework64", "v4.0.30319", "csc.exe"),
          os.path.join(windir, "Microsoft.NET", "Framework", "v4.0.30319", "csc.exe")):
    if os.path.isfile(c):
        CSC_EXE = c
        break

# ===========================================================================
# 1. Documents
# ===========================================================================
def make_documents():
    try:
        p = os.path.join(ROOT, D_DOC, "sample.txt"); ensure_dir(p)
        with open(p, "w", encoding="utf-8") as f: f.write("Plain text sample.\n")
        ok("sample.txt")
    except Exception as e: fail("sample.txt", e)

    try:
        p = os.path.join(ROOT, D_DOC, "sample.md"); ensure_dir(p)
        with open(p, "w", encoding="utf-8") as f: f.write("# Markdown Sample\n")
        ok("sample.md")
    except Exception as e: fail("sample.md", e)

    try:
        p = os.path.join(ROOT, D_DOC, "sample.rtf"); ensure_dir(p)
        with open(p, "w", encoding="ascii", errors="replace") as f:
            f.write(r"{\rtf1\ansi\deff0{\fonttbl{\f0 Arial;}}\fs24 RTF Sample\par}")
        ok("sample.rtf")
    except Exception as e: fail("sample.rtf", e)

    try:
        p = os.path.join(ROOT, D_DOC, "sample.log"); ensure_dir(p)
        with open(p, "w", encoding="utf-8") as f: f.write("[INFO] Sample log entry.\n")
        ok("sample.log")
    except Exception as e: fail("sample.log", e)

    try:
        p = os.path.join(ROOT, D_DOC, "sample.tex"); ensure_dir(p)
        with open(p, "w", encoding="utf-8") as f: f.write("\\documentclass{article}\n\\begin{document}\nHello\n\\end{document}\n")
        ok("sample.tex")
    except Exception as e: fail("sample.tex", e)

    try:
        p = os.path.join(ROOT, D_DOC, "sample.srt"); ensure_dir(p)
        with open(p, "w", encoding="utf-8") as f: f.write("1\n00:00:01,000 --> 00:00:04,000\nSample subtitle\n")
        ok("sample.srt")
    except Exception as e: fail("sample.srt", e)

    try:
        p = os.path.join(ROOT, D_DOC, "sample.vtt"); ensure_dir(p)
        with open(p, "w", encoding="utf-8") as f: f.write("WEBVTT\n\n00:00:01.000 --> 00:00:04.000\nSample caption\n")
        ok("sample.vtt")
    except Exception as e: fail("sample.vtt", e)

    if HAS_ALL_LIBS:
        try:
            p = os.path.join(ROOT, D_DOC, "sample.docx"); ensure_dir(p)
            doc = docx.Document(); doc.add_heading('DOCX Sample', 0); doc.save(p)
            ok("sample.docx")
        except Exception as e: fail("sample.docx", e)

        try:
            p = os.path.join(ROOT, D_DOC, "sample.pdf"); ensure_dir(p)
            c = canvas.Canvas(p, pagesize=letter); c.drawString(100, 750, "PDF Sample"); c.save()
            ok("sample.pdf")
        except Exception as e: fail("sample.pdf", e)

        try:
            p = os.path.join(ROOT, D_DOC, "sample.odt"); ensure_dir(p)
            doc = OpenDocumentText(); doc.text.addElement(H(outlinelevel=1, text="ODT Sample")); doc.save(p)
            ok("sample.odt")
        except Exception as e: fail("sample.odt", e)

        try:
            p = os.path.join(ROOT, D_DOC, "sample.epub"); ensure_dir(p)
            book = epub.EpubBook(); book.set_identifier('id123'); book.set_title('Sample')
            c1 = epub.EpubHtml(title='Ch1', file_name='c1.xhtml'); c1.content = '<h1>Ch1</h1>'
            book.add_item(c1); book.toc = (c1,); book.add_item(epub.EpubNcx()); book.add_item(epub.EpubNav())
            book.spine = ['nav', c1]; epub.write_epub(p, book, {})
            ok("sample.epub")
        except Exception as e: fail("sample.epub", e)

        # Old Office COM (Word)
        try:
            pythoncom.CoInitialize()
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            doc = word.Documents.Add()
            doc.Content.Text = "Sample .doc file generated by Python."
            out = os.path.join(ROOT, D_DOC, "sample.doc"); ensure_dir(out)
            doc.SaveAs(out, FileFormat=0)
            doc.Close()
            ok("sample.doc")
        except Exception as e:
            fail("sample.doc", e)
        finally:
            try: doc.Close()
            except: pass
            try: word.Quit()
            except: pass
            pythoncom.CoUninitialize()

    if CALIBRE_EXE:
        epub_path = os.path.join(ROOT, D_DOC, "sample.epub")
        if os.path.exists(epub_path):
            for ext in ["mobi", "azw3"]:
                out = os.path.join(ROOT, D_DOC, f"sample.{ext}")
                try:
                    subprocess.run([CALIBRE_EXE, epub_path, out], capture_output=True, check=True, timeout=60)
                    ok(f"sample.{ext}")
                except Exception as e: fail(f"sample.{ext}", e)
        else:
            skip("sample.mobi", "sample.epub not found")
            skip("sample.azw3", "sample.epub not found")
    else:
        skip("sample.mobi", "Calibre not found")
        skip("sample.azw3", "Calibre not found")

# ===========================================================================
# 2. Sheets & Presentations
# ===========================================================================
def make_sheets():
    try:
        p = os.path.join(ROOT, D_SHEET, "sample.csv"); ensure_dir(p)
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerows([["Name", "Age"], ["Alice", 28]])
        ok("sample.csv")
    except Exception as e: fail("sample.csv", e)

    try:
        p = os.path.join(ROOT, D_SHEET, "sample.tsv"); ensure_dir(p)
        with open(p, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f, delimiter='\t').writerows([["Name", "Age"], ["Alice", 28]])
        ok("sample.tsv")
    except Exception as e: fail("sample.tsv", e)

    if HAS_ALL_LIBS:
        try:
            p = os.path.join(ROOT, D_SHEET, "sample.xlsx"); ensure_dir(p)
            wb = Workbook(); ws = wb.active; ws.append(["Name", "Age"]); ws.append(["Alice", 28]); wb.save(p)
            ok("sample.xlsx")
        except Exception as e: fail("sample.xlsx", e)

        try:
            p = os.path.join(ROOT, D_SHEET, "sample.pptx"); ensure_dir(p)
            prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = "PPTX Sample"; slide.placeholders[1].text = "Generated by Python"
            prs.save(p); ok("sample.pptx")
        except Exception as e: fail("sample.pptx", e)

        try:
            p = os.path.join(ROOT, D_SHEET, "sample.ods"); ensure_dir(p)
            doc = OpenDocumentSpreadsheet()
            table = Table(name="Sheet1")
            row = TableRow()
            cell = TableCell(valuetype="string")
            cell.addElement(P(text="Name"))
            row.addElement(cell)
            table.addElement(row)
            doc.spreadsheet.addElement(table)
            doc.save(p)
            ok("sample.ods")
        except Exception as e: fail("sample.ods", e)

        try:
            p = os.path.join(ROOT, D_SHEET, "sample.odp"); ensure_dir(p)
            doc = OpenDocumentPresentation()
            page = Page(name="page1", masterpagename="Default")
            frame = Frame(width="22cm", height="3cm", x="2cm", y="2cm")
            textbox = TextBox()
            textbox.addElement(P(text="ODP Sample Slide"))
            frame.addElement(textbox)
            page.addElement(frame)
            doc.presentation.addElement(page)
            doc.save(p)
            ok("sample.odp")
        except Exception as e: fail("sample.odp", e)

        # Old Office COM (Excel)
        try:
            pythoncom.CoInitialize()
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = 0
            wb = excel.Workbooks.Add(); ws = wb.ActiveSheet
            ws.Cells(1, 1).Value = "Name"; ws.Cells(2, 1).Value = "Alice"
            wb.CheckCompatibility = False
            out = os.path.join(ROOT, D_SHEET, "sample.xls"); ensure_dir(out)
            wb.SaveAs(out, FileFormat=56)
            wb.Close()
            ok("sample.xls")
        except Exception as e:
            fail("sample.xls", e)
        finally:
            try: wb.Close()
            except: pass
            try: excel.Quit()
            except: pass
            pythoncom.CoUninitialize()

        # Old Office COM (PowerPoint)
        try:
            pythoncom.CoInitialize()
            ppt = win32com.client.DispatchEx("PowerPoint.Application")
            pres = ppt.Presentations.Add()
            slide = pres.Slides.Add(1, 1)
            slide.Shapes.Title.TextFrame.TextRange.Text = "Sample PPT"
            out = os.path.join(ROOT, D_SHEET, "sample.ppt"); ensure_dir(out)
            pres.SaveAs(out, FileFormat=1)
            pres.Close()
            ok("sample.ppt")
        except Exception as e:
            fail("sample.ppt", e)
        finally:
            try: pres.Close()
            except: pass
            try: ppt.Quit()
            except: pass
            pythoncom.CoUninitialize()

# ===========================================================================
# 3. Images
# ===========================================================================
def make_images():
    if not HAS_ALL_LIBS:
        skip("Images", "Pillow not installed")
        return
    try:
        img = Image.new('RGB', (320, 200), color=(73, 109, 137))
        d = ImageDraw.Draw(img); d.text((100, 90), "Sample", fill=(255, 255, 0))
        for ext, fmt in [("png", "PNG"), ("jpg", "JPEG"), ("jpeg", "JPEG"), ("bmp", "BMP"), ("gif", "GIF"), ("webp", "WEBP"), ("tiff", "TIFF")]:
            p = os.path.join(ROOT, D_IMG, f"sample.{ext}"); ensure_dir(p)
            img.save(p, fmt); ok(f"sample.{ext}")
    except Exception as e: fail("Images", e)

    try:
        p = os.path.join(ROOT, D_IMG, "sample.svg"); ensure_dir(p)
        with open(p, "w") as f:
            f.write('<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100"><rect width="100" height="100" fill="blue"/></svg>')
        ok("sample.svg")
    except Exception as e: fail("sample.svg", e)

    try:
        p = os.path.join(ROOT, D_IMG, "sample.psd"); ensure_dir(p)
        with open(p, "wb") as f:
            f.write(b"8BPS\x00\x01\x00\x00\x00\x00\x00\x00\x00\x03\x00\x00\x00\x40\x00\x00\x00\x40\x00\x08\x00\x03\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00\x00")
        ok("sample.psd")
    except Exception as e: fail("sample.psd", e)

# ===========================================================================
# 4 & 5. Audio and Video
# ===========================================================================
def make_media():
    try:
        p = os.path.join(ROOT, D_AUD, "sample.wav"); ensure_dir(p)
        with wave.open(p, "wb") as w:
            w.setnchannels(1); w.setsampwidth(2); w.setframerate(44100)
            w.writeframes(b"".join(struct.pack("<h", int(10000 * math.sin(2 * math.pi * 440 * i / 44100))) for i in range(44100 * 2)))
        ok("sample.wav")
    except Exception as e: fail("sample.wav", e)

    if not FFMPEG:
        for ext in ["mp3", "aac", "m4a", "flac", "ogg", "wma", "mp4", "avi", "mkv", "mov", "wmv", "flv", "webm"]:
            skip(f"sample.{ext}", "ffmpeg not found")
        return

    for ext in ["mp3", "aac", "m4a", "flac", "ogg", "wma"]:
        out = os.path.join(ROOT, D_AUD, f"sample.{ext}"); ensure_dir(out)
        try:
            subprocess.run([FFMPEG, "-y", "-f", "lavfi", "-i", "sine=frequency=440:duration=2", out], capture_output=True)
            ok(f"sample.{ext}")
        except Exception as e: fail(f"sample.{ext}", e)

    for ext in ["mp4", "avi", "mkv", "mov", "wmv", "flv", "webm"]:
        out = os.path.join(ROOT, D_VID, f"sample.{ext}"); ensure_dir(out)
        try:
            subprocess.run([FFMPEG, "-y", "-f", "lavfi", "-i", "testsrc=size=320x240:rate=15", "-f", "lavfi", "-i", "sine=frequency=440", "-t", "3", "-pix_fmt", "yuv420p", "-shortest", out], capture_output=True)
            ok(f"sample.{ext}")
        except Exception as e: fail(f"sample.{ext}", e)

# ===========================================================================
# 6. Archives
# ===========================================================================
def make_archives():
    demo = b"Sample text for archive.\n"
    for ext, func in [("gz", gzip.compress), ("bz2", bz2.compress), ("xz", lzma.compress)]:
        try:
            p = os.path.join(ROOT, D_ARC, f"sample.{ext}"); ensure_dir(p)
            with open(p, "wb") as f: f.write(func(demo))
            ok(f"sample.{ext}")
        except Exception as e: fail(f"sample.{ext}", e)

    try:
        p = os.path.join(ROOT, D_ARC, "sample.zip"); ensure_dir(p)
        with zipfile.ZipFile(p, "w") as z: z.writestr("readme.txt", demo)
        ok("sample.zip")
    except Exception as e: fail("sample.zip", e)

    try:
        p = os.path.join(ROOT, D_ARC, "sample.tar"); ensure_dir(p)
        with tarfile.open(p, "w") as t:
            info = tarfile.TarInfo("readme.txt")
            info.size = len(demo)
            t.addfile(info, io.BytesIO(demo))
        ok("sample.tar")
    except Exception as e: fail("sample.tar", e)

    try:
        p = os.path.join(ROOT, D_ARC, "sample.jar"); ensure_dir(p)
        with zipfile.ZipFile(p, "w") as z:
            z.writestr("META-INF/MANIFEST.MF", "Manifest-Version: 1.0\nMain-Class: Demo\n\n")
            z.writestr("Demo.class", b"\xca\xfe\xba\xbe")
        ok("sample.jar")
    except Exception as e: fail("sample.jar", e)

    if HAS_ALL_LIBS:
        try:
            p = os.path.join(ROOT, D_ARC, "sample.7z"); ensure_dir(p)
            with py7zr.SevenZipFile(p, 'w') as archive:
                archive.writestr(demo, "readme.txt")
            ok("sample.7z")
        except Exception as e: fail("sample.7z", e)

        try:
            p = os.path.join(ROOT, D_ARC, "sample.iso"); ensure_dir(p)
            iso = pycdlib.PyCdlib(); iso.new(); iso.add_fp(io.BytesIO(demo), len(demo), '/README.TXT;1'); iso.write(p); iso.close()
            ok("sample.iso")
        except Exception as e: fail("sample.iso", e)

    if RAR_EXE:
        try:
            src = os.path.join(ROOT, D_ARC, "_temp.txt")
            with open(src, "wb") as f: f.write(demo)
            out = os.path.join(ROOT, D_ARC, "sample.rar")
            subprocess.run([RAR_EXE, "a", "-y", out, src], capture_output=True)
            ok("sample.rar"); os.remove(src)
        except Exception as e: fail("sample.rar", e)
    else:
        skip("sample.rar", "WinRAR not found")

# ===========================================================================
# 7. Programs
# ===========================================================================
def make_programs():
    try:
        p = os.path.join(ROOT, D_APP, "sample.bat"); ensure_dir(p)
        with open(p, "w") as f: f.write("@echo off\r\necho Hello\r\npause >nul\r\n")
        ok("sample.bat")
    except Exception as e: fail("sample.bat", e)

    try:
        p = os.path.join(ROOT, D_APP, "sample.sh"); ensure_dir(p)
        with open(p, "w") as f: f.write("#!/bin/bash\necho Hello\n")
        ok("sample.sh")
    except Exception as e: fail("sample.sh", e)

    try:
        p = os.path.join(ROOT, D_APP, "sample.ps1"); ensure_dir(p)
        with open(p, "w") as f: f.write("Write-Host 'Hello PowerShell'\n")
        ok("sample.ps1")
    except Exception as e: fail("sample.ps1", e)

    if CSC_EXE:
        try:
            src = os.path.join(ROOT, D_APP, "_temp.cs")
            with open(src, "w") as f: f.write("using System;\nclass P {\n static void Main() {\n Console.WriteLine(\"Hello\");\n }\n}\n")
            out = os.path.join(ROOT, D_APP, "sample.exe")
            subprocess.run([CSC_EXE, "/nologo", "/target:exe", "/out:" + out, src], capture_output=True)
            ok("sample.exe"); os.remove(src)
        except Exception as e: fail("sample.exe", e)

# ===========================================================================
# 8. Code & Data
# ===========================================================================
def make_code_data():
    html_content = """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>HTML Sample</title>
    <style>
        body { font-family: Arial, sans-serif; margin: 40px; }
        .box { border: 1px solid #ccc; padding: 16px; border-radius: 8px; width: 300px; }
        a { display: inline-block; margin: 8px 0; text-decoration: none; color: #0066cc; }
        a:hover { text-decoration: underline; }
    </style>
</head>
<body>
    <h1>HTML Sample Page</h1>
    <div class="box">
        <p>This HTML file includes external search engine links.</p>
        <h3>Search Engines</h3>
        <ul>
            <li><a href="https://www.baidu.com" target="_blank">Baidu (百度)</a></li>
            <li><a href="https://www.bing.com" target="_blank">Bing (必应)</a></li>
        </ul>
    </div>
</body>
</html>
"""
    files = {
        "sample.html": html_content,
        "sample.css": "body { font-family: Arial; }",
        "sample.js": "console.log('Hello');",
        "sample.ts": "const x: number = 1;",
        "sample.py": "print('Hello')",
        "sample.java": "class A { public static void main(String[] a){} }",
        "sample.c": "#include <stdio.h>\nint main(){return 0;}",
        "sample.cpp": "#include <iostream>\nint main(){return 0;}",
        "sample.go": "package main\nfunc main(){}",
        "sample.rs": "fn main(){}",
        "sample.php": "<?php echo 'Hello'; ?>",
        "sample.sql": "SELECT * FROM users;",
        "sample.json": json.dumps({"name": "sample"}),
        "sample.xml": "<root><item>1</item></root>",
        "sample.yaml": "app:\n  name: sample\n",
        "sample.yml": "app:\n  name: sample\n",
        "sample.toml": "[app]\nname = \"sample\"\n",
        "sample.ini": "[Section]\nkey=value\n",
        "sample.properties": "key=value\n",
        "sample.rss": "<rss version=\"2.0\"><channel><title>Sample</title></channel></rss>",
        "sample.xhtml": "<!DOCTYPE html><html><body><h1>XHTML</h1></body></html>"
    }
    for name, content in files.items():
        try:
            p = os.path.join(ROOT, D_CODE, name); ensure_dir(p)
            with open(p, "w", encoding="utf-8") as f: f.write(content)
            ok(name)
        except Exception as e: fail(name, e)

    try:
        p = os.path.join(ROOT, D_CODE, "sample.db"); ensure_dir(p)
        if os.path.exists(p): os.remove(p)
        conn = sqlite3.connect(p); conn.execute("CREATE TABLE t (id INT)"); conn.commit(); conn.close()
        ok("sample.db")
    except Exception as e: fail("sample.db", e)

    try:
        p = os.path.join(ROOT, D_CODE, "sample.sqlite"); ensure_dir(p)
        if os.path.exists(p): os.remove(p)
        shutil.copy(os.path.join(ROOT, D_CODE, "sample.db"), p)
        ok("sample.sqlite")
    except Exception as e: fail("sample.sqlite", e)

# ===========================================================================
# 9. Fonts & 3D/CAD
# ===========================================================================
def make_others():
    # --- TTF (Fixed: Copy system font) ---
    try:
        font_candidates = glob.glob(r"C:\Windows\Fonts\*.ttf")
        target_font = None
        for f in font_candidates:
            if "arial.ttf" in f.lower() or "calibri.ttf" in f.lower():
                target_font = f
                break
        if not target_font and font_candidates:
            target_font = font_candidates[0]
            
        if target_font:
            p = os.path.join(ROOT, D_OTH, "sample.ttf"); ensure_dir(p)
            shutil.copy(target_font, p)
            ok("sample.ttf")
        else:
            skip("sample.ttf", "No system TTF found")
    except Exception as e: fail("sample.ttf", e)

    # --- WOFF (using fontTools to convert the copied TTF) ---
    if HAS_ALL_LIBS:
        try:
            p_ttf = os.path.join(ROOT, D_OTH, "sample.ttf")
            p_woff = os.path.join(ROOT, D_OTH, "sample.woff")
            if os.path.exists(p_ttf):
                font = fontTools.ttLib.TTFont(p_ttf)
                font.flavor = 'woff'
                font.save(p_woff)
                ok("sample.woff")
            else:
                skip("sample.woff", "sample.ttf not found")
        except Exception as e: fail("sample.woff", e)

    # --- DXF (Fixed: Add actual line and circle) ---
    try:
        p = os.path.join(ROOT, D_OTH, "sample.dxf"); ensure_dir(p)
        dxf_content = (
            "0\nSECTION\n2\nENTITIES\n"
            "0\nLINE\n8\n0\n"
            "10\n0.0\n20\n0.0\n"
            "11\n100.0\n21\n100.0\n"
            "0\nCIRCLE\n8\n0\n"
            "10\n50.0\n20\n50.0\n"
            "40\n25.0\n"
            "0\nENDSEC\n0\nEOF\n"
        )
        with open(p, "w", encoding="ascii") as f: f.write(dxf_content)
        ok("sample.dxf")
    except Exception as e: fail("sample.dxf", e)

    # --- OBJ ---
    try:
        p = os.path.join(ROOT, D_OTH, "sample.obj"); ensure_dir(p)
        with open(p, "w") as f: f.write("v 0 0 0\nv 1 0 0\nv 0 1 0\nf 1 2 3\n")
        ok("sample.obj")
    except Exception as e: fail("sample.obj", e)

    # --- STL ---
    try:
        p = os.path.join(ROOT, D_OTH, "sample.stl"); ensure_dir(p)
        with open(p, "w") as f: f.write("solid s\nfacet normal 0 0 -1\nouter loop\nvertex 0 0 0\nvertex 0 1 0\nvertex 1 0 0\nendloop\nendfacet\nendsolid s\n")
        ok("sample.stl")
    except Exception as e: fail("sample.stl", e)

    # --- DAE ---
    try:
        p = os.path.join(ROOT, D_OTH, "sample.dae"); ensure_dir(p)
        with open(p, "w") as f: f.write('<?xml version="1.0" encoding="utf-8"?>\n<COLLADA xmlns="http://www.collada.org/2005/11/COLLADASchema" version="1.4.1"><asset><up_axis>Y_UP</up_axis></asset><library_visual_scenes><visual_scene id="Scene" name="Scene"></visual_scene></library_visual_scenes><scene><instance_visual_scene url="#Scene"/></scene></COLLADA>\n')
        ok("sample.dae")
    except Exception as e: fail("sample.dae", e)

# ===========================================================================
# Main
# ===========================================================================
def main():
    print("=" * 60)
    print("  Ultimate One-Click File Generator (80+ Formats - Complete Final)")
    print(f"  Output: {ROOT}")
    print("=" * 60)

    if os.path.exists(ROOT):
        existing_files = []
        for root, dirs, files in os.walk(ROOT):
            for file in files:
                if file.startswith("sample.") or file.startswith("sample_"):
                    existing_files.append(os.path.join(root, file))
        
        if existing_files:
            print(f"\n[!] Detected {len(existing_files)} existing sample files in target directory:")
            for f in sorted(existing_files):
                print(f"  - {os.path.relpath(f, ROOT)}")
            
            ans = input("\nDo you want to delete them and start fresh? (y/N): ").strip().lower()
            if ans == 'y':
                print("[*] Cleaning up existing sample files...")
                for f in existing_files:
                    try: os.remove(f)
                    except: pass
                print("[*] Cleanup complete. Starting fresh.")
            else:
                print("[*] Keeping existing files. New files will overwrite old ones if names match.")
    else:
        os.makedirs(ROOT)

    if not HAS_ALL_LIBS:
        print(f"\n[!] Missing libraries: {MISSING_LIB}")
        print("Please run: pip install python-docx openpyxl python-pptx Pillow reportlab odfpy EbookLib fonttools pycdlib py7zr pywin32")
        print("Proceeding with basic formats only...\n")

    steps = [
        ("Documents", make_documents),
        ("Sheets/Presentations", make_sheets),
        ("Images", make_images),
        ("Audio/Video", make_media),
        ("Archives", make_archives),
        ("Programs", make_programs),
        ("Code/Data", make_code_data),
        ("Fonts/3D", make_others),
    ]

    for title, fn in steps:
        print(f"[*] Processing: {title} ...")
        try:
            fn()
        except Exception as e:
            print(f"    !! Category {title} error: {e}")

    print("\n" + "=" * 60)
    print("  GENERATION REPORT")
    print("=" * 60)
    
    print(f"\n[+] Successfully Generated Files ({len(OK)}):")
    for name in sorted(OK):
        print(f"  + {name}")
        
    if SKIPPED:
        print(f"\n[-] Skipped Formats ({len(SKIPPED)}):")
        for name, why in SKIPPED:
            print(f"  - {name:<18} {why}")
            
    if FAILED:
        print(f"\n[!] Failed Formats ({len(FAILED)}):")
        for name, err in FAILED:
            print(f"  ! {name:<18} {err}")
            
    print(f"\nAll files have been saved to: {ROOT}")
    print(f"Total planned formats: 80+ (some may be skipped due to missing external tools).")

if __name__ == "__main__":
    main()